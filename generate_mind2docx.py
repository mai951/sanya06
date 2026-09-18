import sys, json, zipfile
from docx import Document

sys.stdout.reconfigure(encoding='utf-8')

z = zipfile.ZipFile('三亚06地块商业改造外装综合幕墙工程.xmind')
data = json.loads(z.read('content.json').decode('utf-8'))
root = data[0]['rootTopic'] if isinstance(data, list) else data['rootTopic']

doc = Document()
doc.add_heading(root['title'], level=1)


def add_node(node, level=2):
    title = (node.get('title') or 'Untitled').strip()
    if title.startswith('细分主题'):
        return
    level = min(level, 4)
    doc.add_heading(title, level=level)
    for child in node.get('children', {}).get('attached', []):
        add_node(child, level + 1)


for child in root.get('children', {}).get('attached', []):
    add_node(child)

doc.add_heading('施工单位重点关注事项', level=2)
points = [
    '项目工期：2026年10月1日开工，2027年4月30日完工，正式开街2027年10月1日（其中临街位置要求2027年1月30日前完成），共计212个日历天。',
    '改造内容：玻璃幕墙、铝板幕墙、外墙涂料、雨蓬、栏杆、地弹门、电动平移门、氟碳漆、铝合金百叶、石材反坎、不锈钢水槽、后置埋件等相关幕墙材料采购及安装。',
    '材料品牌：按限品牌表及封样清单执行，需提前确认供应商资质与到货周期。',
    '合同条款：保修维修、安全生产目标、廉洁合作、工程预结算编制管理规定、三单资料、限价资料等全专业合同附件。',
    '图纸澄清：增补图纸涉及开启扇改固定玻璃、开启扇改感应门，直接影响报价与工期，需在投标前确认。',
    '区域疑问：B区域顶棚是否包含；C区域设备遮挡涉及玻璃量及收口方式不明确、尺寸标识有误、弧形位置外立面表示有误。',
    '质量控制：执行外装淋水检验操作指引、玻璃型材交接铣口标准节点，做好过程记录与验收。',
    '人员配置：参照各阶段人员分配表编制进场计划，确保关键节点人员到位。',
]
for p in points:
    doc.add_paragraph(p, style='List Bullet')

doc.save('招标文件整理.docx')
print('OK: 招标文件整理.docx')